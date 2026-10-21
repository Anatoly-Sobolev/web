# Сборка DOCX-отчётов по лабораторным работам №1–5.
# Оформление повторяет образец отчёта по ЛР (титульный лист ОмГТУ, реферат, содержание,
# введение, основная часть, заключение, список источников).
# Снимки берутся из docs/lab-0N/screenshots/. Чтобы заменить снимок, положите файл
# с тем же именем и запустите сборку повторно:
#     python scripts/build_reports.py
# Содержание собирается в два прохода: DOCX → PDF (LibreOffice) → номера страниц → DOCX.
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SOFFICE = r'C:\Program Files\LibreOffice\program\soffice.exe'
FIGMA_URL = 'https://www.figma.com/design/sCt5ju49i6680hCjosaYHh/'
REPO_URL = 'https://github.com/Anatoly-Sobolev/web'
ACCESS = 'дата обращения: 14.10.2026'

COMMON_SOURCES = [
    'Лабораторные работы по дисциплине «Веб-программирование»: комплексное проектирование и разработка веб-сайта [Электронный ресурс]. – Файл «материалы/Лабораторные_работы_Веб-программирование.pdf» репозитория проекта.',
    'Методика рейтингового контроля по дисциплине «Основы web-технологий». 4 семестр 2026/2027 учебного года. Направление 09.03.04 «Программная инженерия», группа ПИН-252т / преподаватель Е. О. Викулов [Электронный ресурс]. – Омск : ОмГТУ, 2026.',
    'Темы проектов по дисциплине «Основы web-технологий». Тема 12 «Справочник единиц измерения» [Электронный ресурс]. – Файл «материалы/temy.docx» репозитория проекта.',
]


class Report:
    def __init__(self, number, theme, toc_pages=None):
        self.number = number
        self.theme = theme
        self.toc_pages = toc_pages or {}
        self.doc = Document()
        self.figure_no = 0
        self.table_no = 0
        self.listing_no = 0
        self.headings = []
        self._setup()

    # ---------- базовое оформление ----------
    def _setup(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin, section.right_margin = Cm(2), Cm(1)
        section.top_margin, section.bottom_margin = Cm(2), Cm(2)
        section.different_first_page_header_footer = True
        normal = self.doc.styles['Normal']
        normal.font.name = 'Times New Roman'
        normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        self._set_language()
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)
        for name, size in [('Heading 1', 12), ('Heading 2', 12), ('Heading 3', 12)]:
            st = self.doc.styles[name]
            st.font.name = 'Times New Roman'
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = RGBColor(0, 0, 0)
            rpr = st.element.get_or_add_rPr()
            fonts = rpr.find(qn('w:rFonts'))
            if fonts is None:
                fonts = OxmlElement('w:rFonts')
                rpr.append(fonts)
            for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
                if fonts.get(qn(attr)) is not None:
                    del fonts.attrib[qn(attr)]
            for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
                fonts.set(qn(attr), 'Times New Roman')
            st.paragraph_format.line_spacing = 1.15
            st.paragraph_format.keep_with_next = True
        h1 = self.doc.styles['Heading 1'].paragraph_format
        h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h1.space_before, h1.space_after = Pt(0), Pt(12)
        for name in ('Heading 2', 'Heading 3'):
            pf = self.doc.styles[name].paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.first_line_indent = Cm(1.25)
            pf.space_before, pf.space_after = Pt(6), Pt(6)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._field(footer.add_run(), 'PAGE')

    def _set_language(self):
        # Язык правки — русский, чтобы Word не подчёркивал русские слова.
        styles = self.doc.styles.element
        defaults = styles.find(qn('w:docDefaults'))
        targets = []
        if defaults is not None:
            rpr_default = defaults.find(qn('w:rPrDefault'))
            if rpr_default is not None:
                targets.append(rpr_default.find(qn('w:rPr')))
        targets.append(self.doc.styles['Normal'].element.get_or_add_rPr())
        for rpr in targets:
            if rpr is None:
                continue
            for old in rpr.findall(qn('w:lang')):
                rpr.remove(old)
            lang = OxmlElement('w:lang')
            lang.set(qn('w:val'), 'ru-RU')
            lang.set(qn('w:eastAsia'), 'ru-RU')
            lang.set(qn('w:bidi'), 'ar-SA')
            rpr.append(lang)
        settings = self.doc.settings.element
        theme = settings.find(qn('w:themeFontLang'))
        if theme is None:
            theme = OxmlElement('w:themeFontLang')
            settings.append(theme)
        theme.set(qn('w:val'), 'ru-RU')

    @staticmethod
    def _field(run, code):
        begin = OxmlElement('w:fldChar')
        begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = code
        sep = OxmlElement('w:fldChar')
        sep.set(qn('w:fldCharType'), 'separate')
        txt = OxmlElement('w:t')
        txt.text = '1'
        end = OxmlElement('w:fldChar')
        end.set(qn('w:fldCharType'), 'end')
        for el in (begin, instr, sep, txt, end):
            run._r.append(el)

    def p(self, text='', align='justify', indent=True, bold=False, size=None, after=0, keep=False, italic=False):
        par = self.doc.add_paragraph()
        par.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                         'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
        if indent:
            par.paragraph_format.first_line_indent = Cm(1.25)
        par.paragraph_format.space_after = Pt(after)
        if keep:
            par.paragraph_format.keep_with_next = True
        self._runs(par, text, bold, size, italic)
        return par

    def _runs(self, par, text, bold=False, size=None, italic=False):
        # **жирный** и `код` внутри текста
        for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
            if not part:
                continue
            if part.startswith('**'):
                run = par.add_run(part[2:-2]); run.bold = True
            elif part.startswith('`'):
                run = par.add_run(part[1:-1]); run.font.name = 'Courier New'; run.font.size = Pt(11)
            else:
                run = par.add_run(part); run.bold = bold
            if size and not part.startswith('`'):
                run.font.size = Pt(size)
            if italic:
                run.italic = True

    def bullets(self, items):
        for item in items:
            par = self.p('– ' + item)
            par.paragraph_format.first_line_indent = Cm(1.25)

    def link(self, label, url, before='', after=''):
        par = self.p(before)
        hl = OxmlElement('w:hyperlink')
        rel = self.doc.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        hl.set(qn('r:id'), rel)
        run = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1')
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'single')
        rpr.append(color); rpr.append(u); run.append(rpr)
        t = OxmlElement('w:t'); t.text = label; t.set(qn('xml:space'), 'preserve'); run.append(t)
        hl.append(run)
        par._p.append(hl)
        if after:
            par.add_run(after)
        return par

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text, new_page=True):
        par = self.doc.add_paragraph(text.upper(), style='Heading 1')
        par.paragraph_format.page_break_before = new_page
        self.headings.append((1, text.upper()))
        return par

    def h2(self, text):
        self.headings.append((2, text))
        par = self.doc.add_paragraph(text, style='Heading 2')
        par.paragraph_format.page_break_before = True
        return par

    def h3(self, text):
        self.headings.append((3, text))
        return self.doc.add_paragraph(text, style='Heading 3')

    def figure(self, path, caption, width_cm=17.0):
        full = os.path.join(ROOT, path)
        if not os.path.exists(full):
            raise FileNotFoundError(full)
        self.figure_no += 1
        with Image.open(full) as im:
            ratio = im.height / im.width
        width = min(width_cm, 23.0 / ratio)  # не выше ~23 см
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.paragraph_format.space_before = Pt(6)
        par.add_run().add_picture(full, width=Cm(width))
        cap = self.p(f'Рисунок {self.figure_no} – {caption}', align='center', indent=False, after=8)
        return self.figure_no

    def table(self, caption, header, rows, widths=None, size=10):
        self.table_no += 1
        self.p(f'Таблица {self.table_no} – {caption}', align='left', indent=False, keep=True, after=3)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ''
            self._runs(cell.paragraphs[0], h, bold=True, size=size)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ''
                self._runs(cells[i].paragraphs[0], str(v), size=size)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        for row in t.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.line_spacing = 1.0
        self.p('', after=4)
        return self.table_no

    def listing(self, caption, code):
        self.listing_no += 1
        self.p(f'Листинг {self.listing_no} – {caption}', align='left', indent=False, keep=True, after=3)
        lines = code.strip('\n').split('\n')
        for i, line in enumerate(lines):
            par = self.doc.add_paragraph()
            par.paragraph_format.line_spacing = 1.0
            if i < len(lines) - 1:
                par.paragraph_format.keep_with_next = len(lines) < 25
            run = par.add_run(line if line else ' ')
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        self.p('', after=4)
        return self.listing_no

    # ---------- типовые разделы ----------
    def title_page(self, theme_line):
        center = lambda t, size=12, bold=False, after=0: self.p(t, align='center', indent=False, bold=bold, size=size, after=after)
        center('Министерство науки и высшего образования РФ')
        center('Федеральное государственное автономное образовательное учреждение')
        center('высшего образования')
        center('"Омский государственный технический университет"')
        center('Кафедра "Автоматизированные системы обработки информации и управления"', after=48)
        for _ in range(3):
            self.p('')
        center('ОТЧЕТ', size=14, bold=True)
        center(f'ПО ЛАБОРАТОРНОЙ РАБОТЕ №{self.number}', size=14, bold=True, after=6)
        center('по дисциплине «Основы web-технологий»', after=6)
        center(f'«{theme_line}»', after=6)
        center('Проект «Мера — справочник единиц измерения»', after=6)
        center('студента Соболева Анатолия Сергеевича группы ПИН-252т', after=24)
        self.p('')
        center('Пояснительная записка')
        center('Направление подготовки 09.03.04 Программная инженерия', after=48)
        for _ in range(4):
            self.p('')
        t = self.doc.add_table(rows=2, cols=2)
        data = [('Преподаватель\nк.т.н., доцент', 'Е. О. Викулов'), ('Студент', 'А. С. Соболев')]
        for r, (a, b) in enumerate(data):
            t.rows[r].cells[0].text = a
            t.rows[r].cells[1].text = b
            t.rows[r].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for c in t.rows[r].cells:
                c.width = Cm(8.5)
        for _ in range(6):
            self.p('')
        center('Омск 2026')

    def abstract(self, summary, keywords, obj, goal, methods, results, scope):
        self.h1('Реферат')
        self.p(summary)
        self.p('')
        self.p(keywords.upper())
        self.p('')
        for label, text in [('Объект исследования', obj), ('Цель работы', goal), ('Методы исследования', methods),
                            ('Результаты работы', results), ('Область применения', scope)]:
            self.p(f'**{label}:** {text}')

    def toc_placeholder(self):
        self.h1('Содержание')
        self._toc_anchor = self.doc.add_paragraph()

    def fill_toc(self):
        anchor = self._toc_anchor
        entries = [h for h in self.headings if h[1] != 'СОДЕРЖАНИЕ']
        from docx.text.paragraph import Paragraph
        for level, text in entries:
            new_p = OxmlElement('w:p')
            anchor._p.addprevious(new_p)
            para = Paragraph(new_p, anchor._parent)
            para.paragraph_format.left_indent = Cm(0.75 * (level - 1))
            para.paragraph_format.tab_stops.add_tab_stop(Cm(18), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            page = self.toc_pages.get(text, '')
            para.add_run(f'{text}\t{page}')
        return entries

    def sources(self, items):
        self.h1('Список использованных источников')
        for i, item in enumerate(items, 1):
            if isinstance(item, tuple):
                text, url = item
                par = self.p(f'{i}. {text} – URL: ', align='left', indent=False)
                hl = OxmlElement('w:hyperlink')
                rel = self.doc.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
                hl.set(qn('r:id'), rel)
                run = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
                c = OxmlElement('w:color'); c.set(qn('w:val'), '0563C1'); rpr.append(c)
                u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rpr.append(u); run.append(rpr)
                t = OxmlElement('w:t'); t.text = url; run.append(t); hl.append(run); par._p.append(hl)
                par.add_run(f' ({ACCESS}).')
            else:
                self.p(f'{i}. {item}', align='left', indent=False)

    def save(self, path):
        self.fill_toc()
        self.doc.save(path)


def remove_empty_toc_paragraphs(docx_path):
    doc = Document(docx_path)
    body = doc.element.body
    seen_toc = False
    for p in list(body.iterchildren(qn('w:p'))):
        text = ''.join(t.text or '' for t in p.iter(qn('w:t')))
        if text == 'СОДЕРЖАНИЕ':
            seen_toc = True
            continue
        if seen_toc:
            if text == '' and not list(p.iter(qn('w:br'))) and not list(p.iter(qn('w:drawing'))):
                body.remove(p)
                continue
            if text and '\t' not in text:
                break
    doc.save(docx_path)


def pdf_pages(docx_path, workdir):
    subprocess.run([SOFFICE, '--headless', '--convert-to', 'pdf', '--outdir', workdir, docx_path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = os.path.join(workdir, os.path.splitext(os.path.basename(docx_path))[0] + '.pdf')
    return pdf, [page.extract_text() or '' for page in PdfReader(pdf).pages]


def locate(headings, pages):
    norm = lambda s: re.sub(r'\s+', ' ', s).strip()
    result = {}
    start = 0
    for i, t in enumerate(pages):
        if 'СОДЕРЖАНИЕ' in t:
            start = i + 1
            break
    cursor = 1
    for level, text in headings:
        if text == 'СОДЕРЖАНИЕ':
            cursor = start
            continue
        target = norm(text)
        for i in range(cursor, len(pages)):
            if target in norm(pages[i]):
                result[text] = i + 1
                cursor = i
                break
    return result


def build(builder, filename):
    out = os.path.join(ROOT, filename)
    work = tempfile.mkdtemp(prefix='report_')
    tmp = os.path.join(work, 'draft.docx')
    report = builder({})
    report.save(tmp)
    remove_empty_toc_paragraphs(tmp)
    _, pages = pdf_pages(tmp, work)
    numbers = locate(report.headings, pages)
    report = builder(numbers)
    report.save(out)
    remove_empty_toc_paragraphs(out)
    pdf, pages = pdf_pages(out, work)
    final = locate(report.headings, pages)
    missing = [h for _, h in report.headings if h != 'СОДЕРЖАНИЕ' and h not in final]
    if final != numbers:
        report = builder(final)
        report.save(out)
        remove_empty_toc_paragraphs(out)
        pdf, pages = pdf_pages(out, work)
    keep = os.environ.get('REPORT_PDF_DIR')
    if keep:
        shutil.copy(pdf, os.path.join(keep, os.path.basename(out).replace('.docx', '.pdf')))
    print(f'{filename}: страниц {len(pages)}, рисунков {report.figure_no}, таблиц {report.table_no}, '
          f'листингов {report.listing_no}, не найдено в PDF: {missing}')
    shutil.rmtree(work, ignore_errors=True)


if __name__ == '__main__':
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import report_content
    only = sys.argv[1:] or ['1', '2', '3', '4', '5']
    for n in only:
        builder, filename = report_content.REPORTS[n]
        build(builder, filename)
